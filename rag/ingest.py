from pathlib import Path

from dotenv import load_dotenv
from docx import Document

from langchain_core.documents import Document as LCDocument
from langchain_openai import OpenAIEmbeddings
from langchain_chroma import Chroma

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
DOCUMENT_PATH = BASE_DIR / "documents" / "1405_iran-irp.docx"
CHROMA_PATH = BASE_DIR / "chroma_db"


# Load Word document
word_document = Document(str(DOCUMENT_PATH))

documents = []


# Extract normal paragraphs
for paragraph in word_document.paragraphs:
    text = paragraph.text.strip()

    if text:
        documents.append(
            LCDocument(
                page_content=text,
                metadata={"source": "1405_iran-irp.docx"}
            )
        )


# Extract complete tables
for table_number, table in enumerate(word_document.tables, start=1):

    rows = []

    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]

        if any(cells):
            rows.append(" | ".join(cells))

    if rows:

        table_text = "\n".join(rows)

        documents.append(
            LCDocument(
                page_content=f"TABLE {table_number}\n{table_text}",
                metadata={
                    "source": "1405_iran-irp.docx",
                    "type": "table",
                    "table_number": table_number,
                }
            )
        )


print(f"Number of extracted documents: {len(documents)}")


# Split normal text into chunks
from langchain_text_splitters import RecursiveCharacterTextSplitter

text_documents = [
    doc for doc in documents
    if doc.metadata.get("type") != "table"
]

tables = [
    doc for doc in documents
    if doc.metadata.get("type") == "table"
]

text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=1000,
    chunk_overlap=200
)

text_chunks = text_splitter.split_documents(text_documents)


# Keep each table together
chunks = text_chunks + tables

print(f"Number of chunks: {len(chunks)}")


# Embeddings
embeddings = OpenAIEmbeddings(
    model="text-embedding-3-small"
)


# Delete old collection
vector_store = Chroma(
    collection_name="actuarial_documents",
    embedding_function=embeddings,
    persist_directory=str(CHROMA_PATH)
)

vector_store.delete_collection()


# Create fresh collection
vector_store = Chroma(
    collection_name="actuarial_documents",
    embedding_function=embeddings,
    persist_directory=str(CHROMA_PATH)
)


# Store documents
vector_store.add_documents(chunks)

print(f"Stored {len(chunks)} chunks in Chroma.")